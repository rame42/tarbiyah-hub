"""Generate representative Arabic/English test documents (original text written for tests).

Run: python tests/fixtures/make_fixtures.py   (tests call build_all() automatically)
"""
from __future__ import annotations

from pathlib import Path

HERE = Path(__file__).parent
OUT = HERE / "generated"

AR_PAGES = [
    "تربية الأطفال في ضوء المنهج الإسلامي\n\n"
    "إن تربية الأطفال على الأخلاق الحميدة من أهم واجبات الوالدين، فالطفل يولد على الفطرة، "
    "والأسرة هي المحضن الأول الذي يتعلم فيه القيم والآداب. ويؤكد علماء التربية أن القدوة الحسنة "
    "أقوى أثرًا من كثرة التوجيه بالكلام.\n\n"
    "مراحل النمو: تختلف حاجات الطفل في مرحلة الطفولة المبكرة عن حاجاته في مرحلة المراهقة، "
    "ولذلك ينبغي أن يراعي المربي خصائص كل مرحلة عند اختيار الأنشطة والقصص المناسبة.",
    "دور المربي والمعلم\n\n"
    "يحتاج المربي إلى كفايات متعددة، منها: حسن الإصغاء، والحوار، والقدرة على تعديل السلوك "
    "بالتعزيز الإيجابي بدل العقاب. وقد ذكر الغزالي في كتابه إحياء علوم الدين أن الصبي أمانة عند والديه.\n\n"
    "توصية: ينبغي إعداد برامج تدريب المربين تجمع بين التأصيل الشرعي والمهارات التربوية الحديثة.",
]

EN_PAGES = [
    "Positive Parenting and Character Education in Muslim Families\n\n"
    "Parents play the central role in the character development of children. Studies of parenting styles "
    "suggest that an authoritative style, combining warmth with clear expectations, is associated with better "
    "emotional regulation in children than authoritarian or permissive styles.\n\n"
    "Parent-child communication: regular dialogue and attentive listening strengthen the parent-child "
    "relationship and support the transmission of values.",
    "Teacher Training and Mentoring\n\n"
    "Effective educators need structured professional development. Training of trainers programs that model "
    "active learning and reflective mentoring help teachers apply new teaching methods in study circles.\n\n"
    "Limitation: most of the evaluated programs were small, and few measured long-term outcomes for adolescents.",
]

AR_HTML = """<!doctype html><html lang="ar" dir="rtl"><head><meta charset="utf-8">
<title>أساليب الحوار الأسري مع الأبناء</title>
<meta name="author" content="مركز الدراسات التربوية"><meta name="description" content="مقال عن الحوار بين الوالدين والأبناء">
<link rel="license" href="https://creativecommons.org/licenses/by/4.0/"></head>
<body><nav>الرئيسية | مقالات</nav><article><h1>أساليب الحوار الأسري مع الأبناء</h1>
<p>الحوار الأسري وسيلة أساسية لبناء الثقة بين الوالدين والأبناء، وهو من أهم أساليب التربية الإيجابية في الأسرة المسلمة.</p>
<p>ومن آداب الحوار: الإصغاء الجيد، واحترام رأي الطفل، واختيار الوقت المناسب للحديث، والبعد عن أسلوب الأوامر المستمرة.</p>
<p>وتشير دراسات حديثة إلى أن الأسر التي تخصص وقتًا يوميًا للحوار يقل فيها الصراع بين الإخوة.</p>
<p><a href="/articles/teacher-training">تدريب المربين</a> <a href="https://example.org/files/guide.pdf">دليل الحوار (PDF)</a></p>
</article><footer>جميع الحقوق محفوظة</footer></body></html>"""

EN_HTML = """<!doctype html><html lang="en"><head><meta charset="utf-8"><title>Storytelling as a Teaching Method for Young Children</title>
<meta name="author" content="Amina Rahman"><meta name="description" content="How storytelling supports moral education in early childhood">
</head><body><header>Site header</header><main><article><h1>Storytelling as a Teaching Method for Young Children</h1>
<p>Storytelling is one of the oldest teaching methods. In early childhood, stories about the prophets and companions
give children concrete models of patience, honesty and kindness.</p>
<p>Teachers in kindergarten classrooms report that short stories followed by a simple activity help children remember
the moral lesson. Parents can use the same approach at home before bedtime.</p>
<p><a href="/blog/page/2">Older posts</a> <a href="/resources/curriculum.pdf">Curriculum outline</a></p>
</article></main></body></html>"""


def _pdf(pages: list[str], path: Path, rtl: bool) -> None:
    import pymupdf

    doc = pymupdf.open()
    arch = pymupdf.Archive("/usr/share/fonts/truetype/dejavu") if Path("/usr/share/fonts/truetype/dejavu").exists() else None
    css = "@font-face {font-family: dv; src: url(DejaVuSans.ttf);} p {font-family: dv; font-size: 11pt;}" if arch else ""
    for text in pages:
        page = doc.new_page()
        html = "".join(f'<p dir="{"rtl" if rtl else "ltr"}">{para}</p>' for para in text.split("\n\n"))
        page.insert_htmlbox(pymupdf.Rect(50, 50, 545, 790), html, css=css, archive=arch)
    doc.set_metadata({"title": pages[0].split("\n")[0], "author": "Test Author"})
    doc.save(path)


def _scanned_pdf(path: Path) -> None:
    import pymupdf

    src = pymupdf.open()
    p = src.new_page()
    p.insert_text((72, 100), "Scanned page: early childhood curriculum", fontsize=14)
    pix = p.get_pixmap(dpi=100)
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(page.rect, pixmap=pix)  # image only: no text layer
    doc.save(path)


def _docx(path: Path) -> None:
    import docx

    d = docx.Document()
    d.core_properties.title = "Murabbi Competency Framework (draft)"
    d.core_properties.author = "Training Committee"
    d.add_heading("Murabbi Competency Framework", level=1)
    d.add_paragraph("A murabbi needs competencies in mentoring, facilitation of halaqa study circles, and character.")
    d.add_heading("كفايات المربي", level=2)
    d.add_paragraph("من كفايات المربي: القدوة الحسنة، والقدرة على التوجيه والإرشاد، وإدارة الحلقات التربوية.")
    d.save(path)


def _epub(path: Path) -> None:
    from ebooklib import epub

    b = epub.EpubBook()
    b.set_identifier("test-epub-1")
    b.set_title("Rights of Children in Islamic Tradition")
    b.set_language("en")
    b.add_author("Test Scholar")
    b.add_metadata("DC", "rights", "Public domain")
    c1 = epub.EpubHtml(title="Chapter 1", file_name="c1.xhtml", lang="en")
    c1.content = "<h1>Chapter 1: Upbringing</h1><p>Classical scholars such as Ibn Qayyim wrote about the rights of the child, including good upbringing, education and kindness.</p>"
    b.add_item(c1)
    b.toc = (c1,)
    b.spine = ["nav", c1]
    b.add_item(epub.EpubNcx())
    b.add_item(epub.EpubNav())
    epub.write_epub(str(path), b)


def build_all() -> Path:
    OUT.mkdir(parents=True, exist_ok=True)
    targets = {
        "ar_tarbiyah.pdf": lambda p: _pdf(AR_PAGES, p, True),
        "en_parenting.pdf": lambda p: _pdf(EN_PAGES, p, False),
        "scanned.pdf": _scanned_pdf,
        "framework.docx": _docx,
        "children_rights.epub": _epub,
    }
    for name, fn in targets.items():
        p = OUT / name
        if not p.exists():
            fn(p)
    (OUT / "ar_article.html").write_text(AR_HTML, encoding="utf-8")
    (OUT / "en_article.html").write_text(EN_HTML, encoding="utf-8")
    return OUT


if __name__ == "__main__":
    print(build_all())
