"""Text + metadata extraction from documents.

Each extractor returns an ExtractResult made of *segments* (page or section) that
keep their location (page number / heading / timestamp) so citations can point
back to the exact place in the source.
"""
from __future__ import annotations

import io
import re
import shutil
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from ..text.lang import detect_language


@dataclass
class Segment:
    text: str
    page: int | None = None
    heading: str | None = None
    timestamp: float | None = None


@dataclass
class ExtractResult:
    segments: list[Segment] = field(default_factory=list)
    title: str | None = None
    authors: list[str] = field(default_factory=list)
    language: str | None = None
    published: str | None = None
    description: str | None = None
    method: str = ""
    page_count: int | None = None
    ocr_pages: list[int] = field(default_factory=list)  # pages that had no text layer
    links: list[dict[str, str]] = field(default_factory=list)
    license_hint: str | None = None
    meta: dict[str, Any] = field(default_factory=dict)

    @property
    def full_text(self) -> str:
        return "\n\n".join(s.text for s in self.segments if s.text)

    @property
    def word_count(self) -> int:
        return len(self.full_text.split())


class ExtractionError(RuntimeError):
    pass


def clean_text(t: str) -> str:
    # NFKC turns Arabic presentation forms (common in PDFs) back into regular letters.
    t = unicodedata.normalize("NFKC", t)
    t = t.replace("­", "").replace("\x00", "")
    t = re.sub(r"[ \t ]+", " ", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()


def text_quality(text: str) -> float:
    """0..1 heuristic: share of 'normal' characters and plausible word lengths.
    Indicative only (flags garbled OCR / broken encodings); not an objective measure."""
    if not text:
        return 0.0
    sample = text[:20000]
    good = sum(1 for c in sample if c.isalnum() or c.isspace() or c in ".,;:!?()[]\"'-–—،؛؟«»/%")
    bad = sample.count("�") + sum(1 for c in sample if unicodedata.category(c) in ("Co", "Cn"))
    ratio = max(0.0, (good - 5 * bad) / len(sample))
    words = sample.split()
    if words:
        avg = sum(len(w) for w in words) / len(words)
        if avg < 2 or avg > 18:
            ratio *= 0.6
    return round(min(1.0, ratio), 3)


# ------------------------------------------------------------------ formats

def extract_pdf(path: Path) -> ExtractResult:
    import pymupdf

    try:
        doc = pymupdf.open(path)
    except Exception as e:  # corrupted or encrypted
        raise ExtractionError(f"Cannot open PDF: {e}") from e
    if doc.needs_pass:
        raise ExtractionError("PDF is password-protected; not processed.")
    res = ExtractResult(method="pymupdf", page_count=doc.page_count)
    md = doc.metadata or {}
    res.title = (md.get("title") or "").strip() or None
    if md.get("author"):
        res.authors = [a.strip() for a in re.split(r"[;,&]| and ", md["author"]) if a.strip()]
    for i, page in enumerate(doc, start=1):
        # no sort=True: re-sorting spans breaks right-to-left (Arabic) word order
        txt = clean_text(page.get_text("text"))
        if len(txt) < 25 and page.get_images():
            res.ocr_pages.append(i)
        res.segments.append(Segment(text=txt, page=i))
        for link in page.get_links():
            if link.get("uri"):
                res.links.append({"url": link["uri"], "text": ""})
    doc.close()
    return res


def ocr_pdf_pages_tesseract(path: Path, pages: list[int], languages: str = "ara+eng",
                            tesseract_cmd: str = "") -> dict[int, str]:
    import pymupdf
    import pytesseract
    from PIL import Image

    if tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = tesseract_cmd
    out: dict[int, str] = {}
    doc = pymupdf.open(path)
    for p in pages:
        pix = doc[p - 1].get_pixmap(dpi=220)
        img = Image.open(io.BytesIO(pix.tobytes("png")))
        out[p] = clean_text(pytesseract.image_to_string(img, lang=languages))
    doc.close()
    return out


def tesseract_available(tesseract_cmd: str = "") -> bool:
    if tesseract_cmd and Path(tesseract_cmd).exists():
        return True
    return shutil.which("tesseract") is not None


def tesseract_languages(tesseract_cmd: str = "") -> list[str]:
    try:
        import pytesseract

        if tesseract_cmd:
            pytesseract.pytesseract.tesseract_cmd = tesseract_cmd
        return list(pytesseract.get_languages(config=""))
    except Exception:
        return []


def extract_html(content: bytes | str, url: str | None = None) -> ExtractResult:
    import trafilatura
    from bs4 import BeautifulSoup

    html = content.decode("utf-8", errors="replace") if isinstance(content, bytes) else content
    res = ExtractResult(method="trafilatura")
    soup = BeautifulSoup(html, "lxml")
    lic = soup.find("a", rel=lambda v: v and "license" in v) or soup.find("link", rel=lambda v: v and "license" in v)
    if lic and lic.get("href"):
        res.license_hint = lic["href"]
    for a in soup.find_all("a", href=True):
        res.links.append({"url": a["href"], "text": a.get_text(" ", strip=True)[:200]})
    meta = trafilatura.extract_metadata(html, default_url=url)
    if meta:
        res.title = meta.title
        if meta.author:
            res.authors = [a.strip() for a in re.split(r"[;|]", meta.author) if a.strip()]
        res.published = meta.date
        res.description = meta.description
        res.meta["sitename"] = meta.sitename
    body = trafilatura.extract(html, url=url, include_comments=False, include_tables=True, favor_recall=True)
    if not body:
        res.method = "bs4"
        for t in soup(["script", "style", "nav", "footer", "header", "noscript"]):
            t.decompose()
        body = soup.get_text("\n", strip=True)
    if not res.title and soup.title:
        res.title = soup.title.get_text(strip=True)
    html_tag = soup.find("html")
    if html_tag and html_tag.get("lang"):
        res.meta["declared_lang"] = html_tag["lang"]
    res.segments = [Segment(text=clean_text(body or ""))]
    return res


def extract_epub(path: Path) -> ExtractResult:
    import ebooklib
    from bs4 import BeautifulSoup
    from ebooklib import epub

    try:
        book = epub.read_epub(str(path), options={"ignore_ncx": True})
    except Exception as e:
        raise ExtractionError(f"Cannot open EPUB: {e}") from e
    res = ExtractResult(method="ebooklib")
    t = book.get_metadata("DC", "title")
    res.title = t[0][0] if t else None
    res.authors = [a[0] for a in book.get_metadata("DC", "creator")]
    lang = book.get_metadata("DC", "language")
    if lang:
        res.meta["declared_lang"] = lang[0][0]
    rights = book.get_metadata("DC", "rights")
    if rights:
        res.license_hint = rights[0][0]
    for i, item in enumerate(book.get_items_of_type(ebooklib.ITEM_DOCUMENT), start=1):
        soup = BeautifulSoup(item.get_content(), "lxml")
        h = soup.find(["h1", "h2", "h3"])
        txt = clean_text(soup.get_text("\n", strip=True))
        if txt:
            res.segments.append(Segment(text=txt, heading=h.get_text(strip=True) if h else None, page=None))
    res.meta["sections"] = len(res.segments)
    return res


def extract_docx(path: Path) -> ExtractResult:
    import docx

    try:
        d = docx.Document(str(path))
    except Exception as e:
        raise ExtractionError(f"Cannot open DOCX: {e}") from e
    res = ExtractResult(method="python-docx")
    cp = d.core_properties
    res.title = cp.title or None
    if cp.author:
        res.authors = [cp.author]
    heading = None
    buf: list[str] = []
    for p in d.paragraphs:
        if p.style is not None and p.style.name.lower().startswith("heading") and p.text.strip():
            if buf:
                res.segments.append(Segment(text=clean_text("\n".join(buf)), heading=heading))
                buf = []
            heading = p.text.strip()
            if not res.title:
                res.title = heading
        elif p.text.strip():
            buf.append(p.text)
    for table in d.tables:
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        buf.append("\n".join(rows))
    if buf:
        res.segments.append(Segment(text=clean_text("\n".join(buf)), heading=heading))
    return res


def decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16"):
        try:
            s = data.decode(enc)
            if enc == "utf-16" and not data[:2] in (b"\xff\xfe", b"\xfe\xff"):
                continue
            return s
        except UnicodeDecodeError:
            continue
    try:
        return data.decode("cp1256")  # common for legacy Arabic Windows text
    except UnicodeDecodeError:
        return data.decode("latin-1", errors="replace")


def extract_text_file(path: Path) -> ExtractResult:
    txt = clean_text(decode_text(path.read_bytes()))
    res = ExtractResult(method="text")
    first = next((ln.strip() for ln in txt.splitlines() if ln.strip()), "")
    res.title = first[:200] or None
    res.segments = [Segment(text=txt)]
    return res


SUPPORTED_EXT = {".pdf", ".epub", ".docx", ".html", ".htm", ".xhtml", ".txt", ".md"}


def extract_file(path: Path, mime: str | None = None) -> ExtractResult:
    ext = path.suffix.lower()
    mime = (mime or "").split(";")[0].lower()
    if ext == ".pdf" or mime == "application/pdf":
        res = extract_pdf(path)
    elif ext == ".epub" or mime == "application/epub+zip":
        res = extract_epub(path)
    elif ext == ".docx":
        res = extract_docx(path)
    elif ext in (".html", ".htm", ".xhtml") or "html" in mime:
        res = extract_html(path.read_bytes())
    elif ext in (".txt", ".md") or mime.startswith("text/"):
        res = extract_text_file(path)
    else:
        raise ExtractionError(f"Unsupported file type: {ext or mime or 'unknown'}")
    finalize(res)
    return res


def finalize(res: ExtractResult) -> ExtractResult:
    res.language = detect_language(res.full_text) if res.full_text.strip() else (res.language or "und")
    return res


# ------------------------------------------------------------------ chunking

def chunk_segments(segments: list[Segment], target: int = 1400, overlap: int = 150) -> list[dict[str, Any]]:
    """Split into ~target-character passages on paragraph boundaries, never across pages."""
    chunks: list[dict[str, Any]] = []
    for seg in segments:
        if not seg.text.strip():
            continue
        paras = [p.strip() for p in re.split(r"\n\s*\n|\n(?=[^\n]{0,80}\n)", seg.text) if p.strip()]
        if not paras:
            continue
        buf = ""
        for p in paras:
            if len(p) > target * 1.5:  # very long paragraph: split on sentences
                parts = re.split(r"(?<=[.!?؟。])\s+", p)
            else:
                parts = [p]
            for part in parts:
                if buf and len(buf) + len(part) + 1 > target:
                    chunks.append({"text": buf.strip(), "page": seg.page, "heading": seg.heading,
                                   "timestamp": seg.timestamp})
                    tail = buf[-overlap:] if overlap else ""
                    cut = tail.find(" ")
                    buf = (tail[cut + 1:] if cut >= 0 else "") + " " + part
                else:
                    buf = f"{buf}\n{part}" if buf else part
        if buf.strip():
            chunks.append({"text": buf.strip(), "page": seg.page, "heading": seg.heading, "timestamp": seg.timestamp})
    return chunks
